# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl>=3.1.5,<4"]
# ///
"""Read workbook metadata or a selected rectangle; never edit or calculate."""

from __future__ import annotations

import argparse
from contextlib import ExitStack
import json
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries


def inspect(path: Path, sheet: str | None = None, region: str | None = None) -> dict:
    if bool(sheet) != bool(region):
        raise ValueError("Supply both --sheet and --range to inspect cells")
    bounds = None
    if region:
        bounds = range_boundaries(region)
        if any(value is None or value < 1 for value in bounds):
            raise ValueError("Use a finite rectangle such as A1:F20")
        left, top, right, bottom = bounds
        if left > right or top > bottom:
            raise ValueError("Range end must follow its start")

    with ExitStack() as stack:
        wb = load_workbook(path, read_only=True, data_only=False, keep_links=True)
        stack.callback(wb.close)
        result = {
            "path": str(path.resolve()),
            "date_epoch": wb.epoch.isoformat(),
            "calculation_performed": False,
            "notes": [
                "Cached formula results may be missing, stale, or placeholders.",
            ],
        }
        if bounds is not None:
            formulas = wb[sheet]
            if not hasattr(formulas, "iter_rows"):
                raise ValueError(f"{sheet!r} is not a worksheet with cells")
            cached = load_workbook(
                path, read_only=True, data_only=True, keep_links=True
            )
            stack.callback(cached.close)
            values = cached[sheet]
            left, top, right, bottom = bounds
            options = dict(min_col=left, min_row=top, max_col=right, max_row=bottom)
            cells = []
            for row_number, (formula_row, value_row) in enumerate(
                zip(
                    formulas.iter_rows(**options),
                    values.iter_rows(**options),
                    strict=True,
                ),
                start=top,
            ):
                for column_number, (cell, cache) in enumerate(
                    zip(formula_row, value_row, strict=True), start=left
                ):
                    value = cell.value
                    if cell.data_type == "f" and not isinstance(value, str):
                        value = {
                            "kind": type(value).__name__,
                            "text": getattr(value, "text", None),
                            "range": getattr(value, "ref", None),
                        }
                    item = {
                        "cell": f"{get_column_letter(column_number)}{row_number}",
                        "type": cell.data_type,
                        "value": value,
                        "number_format": cell.number_format,
                    }
                    if cell.data_type == "f":
                        item["cached_value"] = cache.value
                        item["cached_type"] = cache.data_type
                    cells.append(item)
            result["selection"] = {"sheet": sheet, "range": region, "cells": cells}
            return result

        # openpyxl owns workbook parsing; ZIP names only provide feature hints.
        archive = stack.enter_context(ZipFile(path))
        parts = archive.namelist()
        feature_prefixes = {
            "charts": "xl/charts/",
            "tables": "xl/tables/",
            "pivot_tables": "xl/pivotTables/",
            "pivot_caches": "xl/pivotCache/",
            "external_links": "xl/externalLinks/",
            "query_tables": "xl/queryTables/",
            "drawings": "xl/drawings/",
            "media": "xl/media/",
        }
        features = {
            name: any(part.startswith(prefix) for part in parts)
            for name, prefix in feature_prefixes.items()
        }
        features["connections"] = "xl/connections.xml" in parts
        features["vba"] = "xl/vbaProject.bin" in parts
        result["sheets"] = [
            {
                "name": ws.title,
                "state": ws.sheet_state,
                "kind": "worksheet" if hasattr(ws, "iter_rows") else "chartsheet",
                "reported_rows": getattr(ws, "max_row", None),
                "reported_columns": getattr(ws, "max_column", None),
            }
            for ws in (wb[name] for name in wb.sheetnames)
        ]
        result["defined_names"] = [
            {"name": name.name, "reference": name.attr_text, "scope": scope}
            for scope, names in [
                ("workbook", wb.defined_names),
                *((ws.title, ws.defined_names) for ws in wb.worksheets),
            ]
            for name in names.values()
        ]
        result["package_features_present"] = features
        result["notes"].extend(
            [
                "Reported dimensions may include formatting or be inaccurate.",
                "Feature presence does not establish round-trip support or validity.",
            ]
        )
        return result


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("--sheet")
    parser.add_argument("--range", dest="region")
    args = parser.parse_args()
    result = inspect(args.input, args.sheet, args.region)
    print(json.dumps(result, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
