# /// script
# requires-python = ">=3.10"
# dependencies = ["python-docx>=1.2,<2"]
# ///
"""Demonstrate DOCX elements: styles, CJK fonts, a table, and a PAGE field."""
from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm


def create(output: Path, latin_font=None, cjk_font=None):
    if output.exists():
        raise FileExistsError(f"Choose a new output: {output}")
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    for name in ("Normal", "Title", "Heading 1"):
        style = doc.styles[name]
        if latin_font:
            style.font.name = latin_font
        if cjk_font:
            fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
            fonts.set(qn("w:eastAsia"), cjk_font)
    header = section.header.paragraphs[0]
    header.text = "页眉示例"
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    field.set(qn("w:dirty"), "true")
    footer._p.append(field)

    doc.add_heading("文档元素示例", 0)
    doc.add_paragraph("中文与 Latin 字体示例。")
    doc.add_heading("表格", 1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    usable_width = section.page_width - section.left_margin - section.right_margin
    widths = [int(usable_width / 3)]
    widths.append(usable_width - sum(widths))
    for column, width in zip(table.columns, widths):
        column.width = width
    for cell, text in zip(table.rows[0].cells, ["元素", "示例用途"]):
        cell.text = text
        for run in cell.paragraphs[0].runs:
            run.bold = True
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in [("表头", "跨页时重复显示"),
                ("列宽", "按页面可用宽度分配")]:
        for cell, text in zip(table.add_row().cells, row):
            cell.text = text
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width
    doc.add_heading("列表", 1)
    for text in ["第一项", "第二项"]:
        doc.add_paragraph(text, style="List Bullet")
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output", type=Path)
    parser.add_argument("--latin-font")
    parser.add_argument("--cjk-font")
    args = parser.parse_args()
    print(create(args.output, args.latin_font, args.cjk_font).resolve())


if __name__ == "__main__":
    main()
